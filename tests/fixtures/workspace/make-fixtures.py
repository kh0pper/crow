#!/usr/bin/env python3
"""Generate the W2 OOXML test fixtures (dev-time only; python-docx/openpyxl/python-pptx, MIT).
Run: python3 tests/fixtures/workspace/make-fixtures.py  (writes rich.docx/.xlsx/.pptx next to this file)."""
import os, io, struct, zlib
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))

def png_1x1():
    raw = b"\x00\xff\x00\x00"
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")

def add_bookmark(par, name, bid):
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(bid)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bid))
    runs = par._p.findall(qn("w:r"))
    runs[0].addnext(start); runs[1].addnext(end)

def docx():
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Casa Nueva — header"
    d.sections[0].footer.paragraphs[0].text = "Página footer"
    d.add_heading("Recetas de la semana", level=1)
    p = d.add_paragraph()
    p.add_run("Tacos al ").bold = True
    p.add_run("pastor").italic = True
    p.add_run(" con piña y jalapeño.")
    add_bookmark(p, "tacos", 1)
    d.add_comment(p.runs[2], text="¿Con salsa verde?", author="Alex", initials="D")  # python-docx >= 1.2
    d.add_heading("Ingredientes", level=2)
    for item in ["Tortillas", "Cebolla", "Cilantro"]:
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("Pasos", level=2)
    for step in ["Marinar la carne", "Asar", "Servir"]:
        d.add_paragraph(step, style="List Number")
    t = d.add_table(rows=2, cols=2); t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "Día", "Plato"
    t.cell(1, 0).text, t.cell(1, 1).text = "Jueves", "Tacos"
    pt = d.add_paragraph("Tab\tseparated text and a line"); pt.add_run().add_break(); pt.add_run("break here.")
    d.add_heading("Notas", level=1)
    d.add_paragraph("Última línea antes del final.")
    d.add_picture(io.BytesIO(png_1x1()))
    d.save(os.path.join(HERE, "rich.docx"))

def xlsx():
    wb = Workbook(); ws = wb.active; ws.title = "Recetas"
    ws.append(["Nombre", "Porciones", "Costo", "Fecha", "Total"])
    rows = [("Tacos", 4, 12.5), ("Pozole", 6, 20), ("Enchiladas", 4, 15.25)]
    for i, (n, s, c) in enumerate(rows, start=2):
        ws.append([n, s, c, None, None])
        ws.cell(row=i, column=4).value = "2026-10-0%d" % i
        ws.cell(row=i, column=5).value = "=B%d*C%d" % (i, i)
    ws["C2"].number_format = "#,##0.00"; ws["B2"].font = Font(bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.merge_cells("A6:C6"); ws["A6"] = "Merged note"
    ws.freeze_panes = "A2"
    chart = BarChart(); chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=4), titles_from_data=True); ws.add_chart(chart, "G2")
    ws2 = wb.create_sheet("Menú semanal"); ws2["A1"] = "Jueves"; ws2["B1"] = "='Recetas'!A2"
    wb.save(os.path.join(HERE, "rich.xlsx"))

def pptx():
    pr = Presentation()
    s1 = pr.slides.add_slide(pr.slide_layouts[0]); s1.shapes.title.text = "Menú de octubre"; s1.placeholders[1].text = "Casa Nueva"
    s2 = pr.slides.add_slide(pr.slide_layouts[1]); s2.shapes.title.text = "Jueves"; s2.placeholders[1].text = "Tacos al pastor\nAgua de jamaica"
    s2.notes_slide.notes_text_frame.text = "Recordar comprar piña"
    tb = s2.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)); tb.text_frame.text = "Texto libre"
    pr.save(os.path.join(HERE, "rich.pptx"))

docx(); xlsx(); pptx(); print("ok")
