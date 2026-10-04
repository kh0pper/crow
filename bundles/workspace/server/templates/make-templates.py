#!/usr/bin/env python3
"""Blank new-file templates for ws_*_create (python-docx/openpyxl/python-pptx, MIT). Dev-time; output committed."""
import os
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__))
d = Document()
for s in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6", "List Bullet", "List Number", "Table Grid"):
    d.styles[s]  # KeyError if the default template ever loses one
body = d.element.body
for p in list(body.iterchildren()):
    if not p.tag.endswith("sectPr"):
        body.remove(p)
d.save(os.path.join(HERE, "blank.docx"))
wb = Workbook(); wb.active.title = "Sheet1"; wb.save(os.path.join(HERE, "blank.xlsx"))
pr = Presentation(); _ = pr.notes_master  # python-pptx creates the notes master lazily: force it so edit_notes works on new decks
pr.save(os.path.join(HERE, "blank.pptx"))
print("ok")
